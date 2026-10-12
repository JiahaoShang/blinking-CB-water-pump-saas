"""企业知识第二批 · 真实文件解析与资源保护（任务2）。

- 解析器（隔离环境依赖，见 requirements-knowledge.txt）：
    pypdf==6.20.0       文字 PDF 按页提取（记录真实页码）
    python-docx==1.2.0  DOCX 按段落 / 标题路径 / 表格行列提取（不虚构页号）
  两者惰性导入：未安装时对应格式明确失败（不静默猜测内容）；本模块可被无解析依赖的
  进程安全导入（第一批模块不加载 pypdf / docx）。
- extract_document(filename, data) 返回：
    status      extracted（有可核对正文段落）/ registered（仅登记原件）/ failed（明确失败）
    paragraphs  [{text, location{...}, source_kind}]，location 只记录可确定的位置：
                PDF 页码；DOCX 段落序号（含标题路径）/ 表格行列；文本行号；人工撰写。
    warnings    如实说明（扫描件、旧 DOC、合并单元格、截断等）
    error       失败原因（failed 时非空）
    format / title_hint（显示辅助；不代表已识别）
- 资源保护：文档 10MiB / 文本 1MiB 在批次创建时校验（本模块仅解析已保存的字节）；
  文字 PDF 最多 200 页；DOCX 解压前检查 ZIP 条目声明大小与压缩比（解压异常拒绝）；
  DOCX 段落数上限截断并给出警告。
- 设计依据：技术计划任务2、PRD v0.3 §9（支持登记不等于识别）、v0.2 §6.2。
"""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

# ── 处理保护上限 ──────────────────────────────────────────────────────────────
PDF_MAX_PAGES = 200                              # 文字 PDF 最多处理页数（超出仍可原件归档/人工整理）
DOCX_MAX_PARAGRAPHS = 5000                       # DOCX 单个文档最多收录段落（超出截断并警告）
DOCX_MAX_UNCOMPRESSED_TOTAL = 64 * 1024 * 1024   # 解压后总体积上限（防解压炸弹）
DOCX_MAX_ENTRY_UNCOMPRESSED = 32 * 1024 * 1024   # 单条目解压后上限
DOCX_SUSPICIOUS_RATIO = 200                      # 压缩比阈值（解压/压缩）
DOCX_SUSPICIOUS_MIN_SIZE = 8 * 1024 * 1024       # 触发压缩比检查的最小解压大小

# ── 格式分组（扩展名 → 处理方式）──────────────────────────────────────────────
TEXT_EXTENSIONS = (".txt", ".md", ".markdown", ".csv")
PDF_EXTENSIONS = (".pdf",)
DOCX_EXTENSIONS = (".docx",)
OLD_DOC_EXTENSIONS = (".doc",)
IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".gif", ".webp", ".heic", ".bmp", ".tif", ".tiff")
VIDEO_EXTENSIONS = (".mp4", ".mov", ".avi", ".mkv", ".webm", ".m4v")
REGISTER_ONLY_EXTENSIONS = OLD_DOC_EXTENSIONS + IMAGE_EXTENSIONS + VIDEO_EXTENSIONS

_FORMAT_BY_EXTENSION = {
    ".txt": "txt", ".md": "md", ".markdown": "md", ".csv": "csv",
    ".pdf": "pdf", ".docx": "docx", ".doc": "doc",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".gif": "image",
    ".webp": "image", ".heic": "image", ".bmp": "image", ".tif": "image", ".tiff": "image",
    ".mp4": "video", ".mov": "video", ".avi": "video", ".mkv": "video", ".webm": "video", ".m4v": "video",
}

FACE_TEXT_CHARS_RE = re.compile(r"[\u4e00-\u9fff]")


def format_for_filename(filename: str) -> str:
    return _FORMAT_BY_EXTENSION.get(Path(str(filename or "")).suffix.lower(), "other")


def decode_text_bytes(data: bytes) -> str:
    """文本解码（UTF-8 优先、兼容 GB18030）；二进制内容（NUL）拒绝。"""
    if b"\x00" in data:
        raise ValueError("文件不是纯文本（包含二进制内容）")
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("文件内容不是可识别的文本编码（支持 UTF-8 / GB）")


def _result(status, fmt, *, paragraphs=None, warnings=None, error="", title_hint="", page_count=None):
    payload = {
        "status": status,
        "format": fmt,
        "title_hint": title_hint,
        "paragraphs": paragraphs or [],
        "warnings": warnings or [],
        "error": error,
    }
    if page_count is not None:
        payload["page_count"] = page_count
    return payload


def _text_paragraphs(text: str) -> list[dict]:
    paragraphs = []
    for index, line in enumerate(text.splitlines(), start=1):
        stripped = line.strip()
        if not stripped:
            continue
        paragraphs.append({
            "text": stripped,
            "location": {"kind": "line", "line": index},
            "source_kind": "text",
        })
    return paragraphs


def _extract_pdf(data: bytes) -> dict:
    try:
        import pypdf
    except ImportError:
        return _result("failed", "pdf", error="解析组件未安装（pypdf）：请在本机隔离环境安装后重试，或人工整理正文。")
    # 损坏文件是预期失败路径：解析已知损坏文件时压低 pypdf 的告警噪音（失败原因由本函数如实给出）。
    import logging

    pdf_logger = logging.getLogger("pypdf")
    previous_level = pdf_logger.level
    pdf_logger.setLevel(logging.ERROR)
    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            return _result("failed", "pdf", error="文件已加密（受密码保护）：请使用未加密文件，或保留原件人工整理正文。")
        page_count = len(reader.pages)
    except Exception:
        return _result("failed", "pdf", error="文件损坏或不是有效的 PDF（无法解析）：请重新导出或人工整理正文。")
    finally:
        pdf_logger.setLevel(previous_level)
    if page_count > PDF_MAX_PAGES:
        return _result(
            "registered", "pdf", page_count=page_count,
            warnings=[f"PDF 共 {page_count} 页，超过 {PDF_MAX_PAGES} 页处理上限：原件已登记，可人工整理正文（不做自动解析）。"])
    paragraphs: list[dict] = []
    warnings: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            warnings.append(f"第 {page_number} 页文字提取失败，请人工核对这一页。")
            continue
        for block in re.split(r"\n\s*\n", page_text):
            block = block.strip()
            if not block:
                continue
            paragraphs.append({
                "text": block,
                "location": {"kind": "page", "page": page_number},
                "source_kind": "pdf_text",
            })
    if not paragraphs:
        warnings.append("没有识别到可提取文字（可能为扫描件或图片型 PDF）：原件已登记，可人工整理正文。")
        return _result("registered", "pdf", warnings=warnings, page_count=page_count)
    return _result("extracted", "pdf", paragraphs=paragraphs, warnings=warnings, page_count=page_count)


def _check_docx_zip_guards(data: bytes):
    """DOCX 解压前的资源保护：返回错误消息或 None。"""
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return "文件损坏或不是有效的 DOCX（无法解压）：请重新导出或人工整理正文。"
    infos = archive.infolist()
    if sum(info.file_size for info in infos) > DOCX_MAX_UNCOMPRESSED_TOTAL:
        return f"解压后体积超过 {DOCX_MAX_UNCOMPRESSED_TOTAL // (1024 * 1024)}MB 保护上限：请拆分文件或人工整理正文。"
    for info in infos:
        if info.file_size > DOCX_MAX_ENTRY_UNCOMPRESSED:
            return "文档内部单项体积异常（超过解压保护上限）：已拒绝自动处理，请人工核对来源或整理正文。"
        if (info.file_size >= DOCX_SUSPICIOUS_MIN_SIZE and info.compress_size > 0
                and info.file_size / info.compress_size >= DOCX_SUSPICIOUS_RATIO):
            return "文档压缩比异常（疑似解压炸弹）：已拒绝自动解压，请提供正常来源文件或人工整理正文。"
    return None


def _extract_docx(data: bytes) -> dict:
    if data[:4] == b"\xd0\xcf\x11\xe0":
        return _result("failed", "docx", error="文件已加密（受密码保护）或为旧式复合格式：请使用未加密的 .docx，或保留原件人工整理正文。")
    guard_error = _check_docx_zip_guards(data)
    if guard_error:
        return _result("failed", "docx", error=guard_error)
    try:
        import docx
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        return _result("failed", "docx", error="解析组件未安装（python-docx）：请在本机隔离环境安装后重试，或人工整理正文。")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        return _result("failed", "docx", error="文件损坏或不是有效的 Word 文档（缺少文档结构）：请重新导出或人工整理正文。")

    paragraphs: list[dict] = []
    warnings: list[str] = []
    heading_stack: list[tuple[int, str]] = []
    paragraph_number = 0
    table_number = 0

    def section_path() -> str:
        return " / ".join(title for _, title in heading_stack)

    def add_paragraph_record(text: str, source_kind: str, heading_level: int | None = None):
        nonlocal paragraph_number
        paragraph_number += 1
        location = {"kind": "paragraph", "paragraph": paragraph_number}
        path = section_path()
        if path:
            location["section"] = path
        paragraphs.append({"text": text, "location": location, "source_kind": source_kind})
        if heading_level is not None:
            while heading_stack and heading_stack[-1][0] >= heading_level:
                heading_stack.pop()
            heading_stack.append((heading_level, text))

    def heading_level(paragraph) -> int | None:
        name = (paragraph.style.name or "") if paragraph.style is not None else ""
        match = re.match(r"^(?:Heading|标题)\s*(\d+)", name)
        if match:
            return int(match.group(1))
        return None

    def iter_blocks():
        body = document.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, document)
            elif child.tag == qn("w:tbl"):
                yield Table(child, document)

    for block in iter_blocks():
        if isinstance(block, Paragraph):
            text = (block.text or "").strip()
            if not text:
                continue
            level = heading_level(block)
            if level is not None:
                add_paragraph_record(text, "docx_heading", heading_level=level)
            else:
                add_paragraph_record(text, "docx_paragraph")
        else:
            table_number += 1
            merged = False
            for row_index, row in enumerate(block.rows, start=1):
                seen_cells = set()
                for col_index, cell in enumerate(row.cells, start=1):
                    cell_key = id(cell._tc)
                    if cell_key in seen_cells:
                        merged = True
                        continue
                    seen_cells.add(cell_key)
                    cell_text = " ".join(part.strip() for part in (cell.text or "").splitlines() if part.strip())
                    if not cell_text:
                        continue
                    paragraphs.append({
                        "text": cell_text,
                        "location": {"kind": "table", "table": table_number, "row": row_index, "col": col_index},
                        "source_kind": "docx_table",
                    })
            if merged:
                warnings.append(f"表格 {table_number} 存在合并单元格：已按可确定的位置列出，请人工核对。")

    if len(paragraphs) > DOCX_MAX_PARAGRAPHS:
        paragraphs = paragraphs[:DOCX_MAX_PARAGRAPHS]
        warnings.append(f"文档段落过多：仅收录前 {DOCX_MAX_PARAGRAPHS} 段，请在原件中人工核对其余内容。")
    if not paragraphs:
        warnings.append("没有识别到可提取文字：原件已登记，可人工整理正文。")
        return _result("registered", "docx", warnings=warnings)
    return _result("extracted", "docx", paragraphs=paragraphs, warnings=warnings)


def extract_document(filename: str, data: bytes) -> dict:
    """对已保存的私有原件字节做真实文字提取（不读取用户机器路径）。

    - 返回 status: extracted / registered / failed（见模块说明）；段落保留可确定位置。
    - 支持文字文件（TXT/MD/CSV）、文字 PDF（pypdf，页码）、DOCX（python-docx，段落/表格）。
    - 图片 / 视频 / 旧版 DOC 仅登记原件；未知格式明确失败；加密/损坏/解压异常明确失败。
    """
    name = str(filename or "").strip()
    if not isinstance(data, (bytes, bytearray)):
        return _result("failed", "other", error="文件内容无效。")
    payload = bytes(data)
    if not payload:
        return _result("failed", format_for_filename(name), error="文件内容为空。")
    extension = Path(name).suffix.lower()
    title_hint = Path(name).stem.strip() or name
    if extension in TEXT_EXTENSIONS:
        fmt = _FORMAT_BY_EXTENSION[extension]
        try:
            text = decode_text_bytes(payload)
        except ValueError as error:
            return _result("failed", fmt, error=str(error), title_hint=title_hint)
        paragraphs = _text_paragraphs(text)
        if not paragraphs:
            return _result("registered", fmt, warnings=["文件没有可核对文字内容：已登记，可人工整理正文。"], title_hint=title_hint)
        return _result("extracted", fmt, paragraphs=paragraphs, title_hint=title_hint)
    if extension in PDF_EXTENSIONS:
        result = _extract_pdf(payload)
        result["title_hint"] = title_hint
        return result
    if extension in DOCX_EXTENSIONS:
        result = _extract_docx(payload)
        result["title_hint"] = title_hint
        return result
    if extension in OLD_DOC_EXTENSIONS:
        return _result("registered", "doc",
                       warnings=["旧版 DOC 不自动解析：原件已登记，可人工整理正文。"], title_hint=title_hint)
    if extension in IMAGE_EXTENSIONS:
        return _result("registered", "image",
                       warnings=["图片仅登记原件，不做文字识别（本批无 OCR）：可人工整理正文。"], title_hint=title_hint)
    if extension in VIDEO_EXTENSIONS:
        return _result("registered", "video",
                       warnings=["视频仅登记原件，不做内容理解：可人工整理正文。"], title_hint=title_hint)
    return _result("failed", format_for_filename(name), title_hint=title_hint,
                   error="暂不支持该格式：请使用 TXT / Markdown / CSV / PDF / DOCX；图片、视频与旧版 DOC 可仅登记原件。")
