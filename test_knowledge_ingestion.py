"""企业知识第二批 · 任务2：真实统一入库、识别与核对（临时 SQLite）。

对应用户已授权技术计划
（docs/superpowers/plans/2026-10-12-enterprise-knowledge-phase2.md 任务2）：
先写有意义的失败测试、再最小实现。覆盖：
- 真实文字 TXT / PDF / DOCX 解析（PDF 页码、DOCX 段落/表格位置，不虚构页号）；
- 扫描型（无可提取文字）/加密/损坏/解压资源异常/超限的明确失败，不把登记说成解析成功；
- 多文件与链接/手动混合批次真实入库；完成识别只创建草稿，绝不自动发布；
- 任务持久化与重开恢复：崩溃中断后可继续、失败可重试且不重复创建；
- 核对 / 编辑（撤销核对并记人工修改）/ 排除；发布闸门（未核对、空正文拦截）；
- 仅登记原件可归档发布（AI 不可用、不算发布知识正文）；人工补充正文需核对；
- 建议已有/新主题（真实文字命中）、改名 / 合并 / 拒绝 / 整批共同主题、确认前不建立、
  同批同名合并、批次确认与发布幂等；
- 字段来源对应真实 source/version/location；权限与跨企业拒绝；原件私有 token 命名 + sha256。

运行环境：隔离虚拟环境（需 pypdf==6.20.0 / python-docx==1.2.0）：
    .venv-knowledge/bin/python -m unittest test_knowledge_ingestion -v
只用临时数据库与测试私有上传目录；不触碰 data/ 下真实演示数据与任何密钥。
"""
import io
import json
import os
import re
import tempfile
import unittest
import zipfile
from pathlib import Path

from app import Database
import knowledge_ingestion

try:
    import pypdf
    import docx as docx_lib

    HAVE_PARSERS = True
except ImportError:  # pragma: no cover - 环境说明
    HAVE_PARSERS = False

REQUIRES_PARSERS = unittest.skipUnless(
    HAVE_PARSERS, "需要隔离环境 .venv-knowledge（pypdf==6.20.0 / python-docx==1.2.0）")

TENANT = "demo-tenant"
OTHER = "other-tenant"


# ── 真实夹具构造（不引用任何外部文件）────────────────────────────────────────

def build_text_pdf(pages):
    """构造真实可解析的文字 PDF（自由文本；仅测试用）。pages: 每页一行文字。"""
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []

    def write(body):
        offsets.append(out.tell())
        out.write(body)
        out.write(b"\n")

    kids = []
    page_bodies = []
    for index, text in enumerate(pages):
        page_index = 4 + 2 * index
        content_index = page_index + 1
        kids.append(f"{page_index} 0 R")
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
        page_bodies.append((
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_index} 0 R >>".encode(),
            b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        ))
    write(b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj")
    write(b"2 0 obj << /Type /Pages /Kids [" + " ".join(kids).encode() + b"] /Count %d >> endobj" % len(pages))
    write(b"3 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj")
    for index, (page, content) in enumerate(page_bodies):
        write(b"%d 0 obj " % (4 + 2 * index) + page + b" endobj")
        write(b"%d 0 obj " % (5 + 2 * index) + content + b" endobj")
    xref_position = out.tell()
    total = 4 + 2 * len(pages)
    out.write(b"xref\n0 %d\n" % total)
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer << /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (total, xref_position))
    return out.getvalue()


def build_encrypted_pdf():
    """真实加密 PDF（无正确密码不可读取）。"""
    reader = pypdf.PdfReader(io.BytesIO(build_text_pdf(["secret content"])))
    writer = pypdf.PdfWriter()
    writer.append(reader)
    writer.encrypt("owner-pass-123")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def build_docx():
    """真实 DOCX：标题 / 段落 / 表格。"""
    document = docx_lib.Document()
    document.add_heading("安装指南", level=1)
    document.add_paragraph("连接电源并检查指示灯。")
    document.add_heading("维护", level=2)
    document.add_paragraph("每月检查滤芯。")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "参数"
    table.rows[0].cells[1].text = "值"
    table.rows[1].cells[0].text = "电压"
    table.rows[1].cells[1].text = "220V"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_zip_bomb():
    """声明解压后 30MB 的高压缩比 ZIP 内容（伪装 docx 的中央目录），用于资源保护测试。"""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\x00" * (30 * 1024 * 1024))
    return buffer.getvalue()


class IngestionTestBase(unittest.TestCase):
    maxDiff = None

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.upload_dir = Path(self.tmp.name) / "uploads"
        os.environ["B2B_KB_STORAGE_ROOT"] = str(self.upload_dir)
        self.db_path = Path(self.tmp.name) / "knowledge2.db"
        self.db = Database(self.db_path)
        self.db.knowledge_storage_root = self.upload_dir
        self.db.seed_demo_users()
        self.admin = self.db.authenticate_user("demo-admin@example.invalid", "demo-admin")
        self.member = self.db.authenticate_user("demo-sales@example.invalid", "demo-sales")
        self.db.ensure_tenant(OTHER, "其他企业（测试）")
        self.other_admin_id = self.db.create_user(
            OTHER, "other-admin@example.invalid", "对方管理员", "admin", "other-admin", "test-setup")

    def tearDown(self):
        self.db.conn.close()
        os.environ.pop("B2B_KB_STORAGE_ROOT", None)
        self.tmp.cleanup()

    def reopen(self):
        self.db.conn.close()
        self.db = Database(self.db_path)
        self.db.knowledge_storage_root = self.upload_dir
        return self.db

    def file_entry(self, filename, data, title=""):
        return {"kind": "file", "filename": filename, "data": data, "title": title}

    def source_by(self, batch, filename):
        for source in batch["sources"]:
            if source["filename"] == filename:
                return source
        raise AssertionError(f"批次中没有来源：{filename}（现有 {[s['filename'] for s in batch['sources']]}）")


# ── extract_document：真实解析与明确失败 ─────────────────────────────────────

@REQUIRES_PARSERS
class ExtractDocumentTests(unittest.TestCase):
    maxDiff = None

    def test_text_file_lines_with_locations(self):
        data = "整机保修 24 个月。\n\n- 覆盖全国\n- 响应 48 小时\n".encode("utf-8")
        result = knowledge_ingestion.extract_document("保修说明.txt", data)
        self.assertEqual(result["status"], "extracted")
        texts = [p["text"] for p in result["paragraphs"]]
        self.assertEqual(texts, ["整机保修 24 个月。", "- 覆盖全国", "- 响应 48 小时"])
        self.assertEqual(result["paragraphs"][0]["location"]["kind"], "line")
        self.assertEqual(result["paragraphs"][0]["location"]["line"], 1)
        self.assertEqual(result["paragraphs"][2]["location"]["line"], 4)
        for paragraph in result["paragraphs"]:
            self.assertEqual(paragraph["source_kind"], "text")

    def test_text_file_gb18030_decodes(self):
        result = knowledge_ingestion.extract_document("说明.csv", "项目,值\n电压,220V\n".encode("gb18030"))
        self.assertEqual(result["status"], "extracted")
        self.assertIn("电压,220V", result["paragraphs"][1]["text"])

    def test_binary_content_fails_clearly(self):
        result = knowledge_ingestion.extract_document("坏文件.txt", b"abc\x00def")
        self.assertEqual(result["status"], "failed")
        self.assertTrue(result["error"])
        self.assertEqual(result["paragraphs"], [])

    def test_real_pdf_page_locations(self):
        data = build_text_pdf(["Warranty 36 months.", "Service 48 hours."])
        result = knowledge_ingestion.extract_document("保修.pdf", data)
        self.assertEqual(result["status"], "extracted")
        self.assertEqual(len(result["paragraphs"]), 2)
        first, second = result["paragraphs"]
        self.assertIn("Warranty", first["text"])
        self.assertEqual(first["location"]["kind"], "page")
        self.assertEqual(first["location"]["page"], 1)
        self.assertEqual(second["location"]["page"], 2)
        self.assertEqual(first["source_kind"], "pdf_text")
        self.assertEqual(result["page_count"], 2)

    def test_real_docx_paragraph_and_table_locations_without_pages(self):
        result = knowledge_ingestion.extract_document("指南.docx", build_docx())
        self.assertEqual(result["status"], "extracted")
        by_kind = {}
        for paragraph in result["paragraphs"]:
            by_kind.setdefault(paragraph["source_kind"], []).append(paragraph)
        texts = [p["text"] for p in by_kind["docx_paragraph"]]
        self.assertIn("连接电源并检查指示灯。", texts)
        self.assertIn("每月检查滤芯。", texts)
        first = by_kind["docx_paragraph"][0]
        self.assertEqual(first["location"]["kind"], "paragraph")
        self.assertEqual(first["location"]["section"], "安装指南")
        maintenance = [p for p in by_kind["docx_paragraph"] if p["text"] == "每月检查滤芯。"][0]
        self.assertEqual(maintenance["location"]["section"], "安装指南 / 维护")
        cell = [p for p in by_kind["docx_table"] if p["text"] == "电压"][0]
        self.assertEqual(cell["location"]["kind"], "table")
        self.assertEqual(cell["location"]["table"], 1)
        self.assertEqual((cell["location"]["row"], cell["location"]["col"]), (2, 1))
        # 不虚构页号：位置里没有 page 字段。
        for paragraph in result["paragraphs"]:
            self.assertNotIn("page", paragraph["location"])

    def test_blank_pdf_registered_not_parsed(self):
        """扫描型 PDF：没有可提取文字 → 仅登记，不显示解析成功。"""
        result = knowledge_ingestion.extract_document("扫描件.pdf", build_text_pdf(["", ""]))
        self.assertEqual(result["status"], "registered")
        self.assertEqual(result["paragraphs"], [])
        self.assertTrue(result["warnings"])
        self.assertIn("文字", " ".join(result["warnings"]))

    def test_encrypted_pdf_failed_with_reason(self):
        result = knowledge_ingestion.extract_document("加密.pdf", build_encrypted_pdf())
        self.assertEqual(result["status"], "failed")
        self.assertIn("加密", result["error"])

    def test_corrupt_pdf_failed_with_reason(self):
        result = knowledge_ingestion.extract_document("损坏.pdf", b"%PDF-1.4\nthis is not a pdf")
        self.assertEqual(result["status"], "failed")
        self.assertTrue(result["error"])
        self.assertNotIn("成功", result["error"])

    def test_corrupt_docx_failed_with_reason(self):
        result = knowledge_ingestion.extract_document("损坏.docx", b"PK\x03\x04not-a-real-docx")
        self.assertEqual(result["status"], "failed")
        self.assertTrue(result["error"])

    def test_zip_bomb_docx_rejected_without_extraction(self):
        """高压缩比 ZIP（解压后 >100 倍且超阈值）直接拒绝：不无限解压。"""
        result = knowledge_ingestion.extract_document("炸弹.docx", build_zip_bomb())
        self.assertEqual(result["status"], "failed")
        self.assertTrue(result["error"])
        self.assertEqual(result["paragraphs"], [])

    def test_pdf_page_limit_protection(self):
        """文字 PDF 最多 200 页处理保护：超出 → 仅登记原件 + 明确说明，可人工整理。"""
        result = knowledge_ingestion.extract_document("超长.pdf", build_text_pdf([""] * 201))
        self.assertEqual(result["status"], "registered")
        self.assertEqual(result["paragraphs"], [])
        self.assertIn("200", " ".join(result["warnings"]))

    def test_unsupported_extension_failed(self):
        result = knowledge_ingestion.extract_document("程序.exe", b"MZ\x90\x00")
        self.assertEqual(result["status"], "failed")
        self.assertIn("支持", result["error"])

    def test_register_only_formats(self):
        old_doc = knowledge_ingestion.extract_document("旧文档.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
        self.assertEqual(old_doc["status"], "registered")
        self.assertTrue(old_doc["warnings"])
        image = knowledge_ingestion.extract_document("照片.png", b"\x89PNG\r\n\x1a\n" + b"0" * 32)
        self.assertEqual(image["status"], "registered")
        self.assertEqual(image["paragraphs"], [])


# ── 批次服务：真实入库、恢复、发布、建议确认 ─────────────────────────────────

@REQUIRES_PARSERS
class BatchServiceTests(IngestionTestBase):

    def create_mixed_batch(self):
        entries = [
            self.file_entry("保修说明.pdf", build_text_pdf(["Warranty 36 months.", "Service 48 hours."])),
            self.file_entry("安装指南.docx", build_docx()),
            {"kind": "link", "url": "https://example.invalid/manual", "title": "在线手册"},
            {"kind": "manual", "text": "整机保修政策适用于全部机型。\n支持以换代修。", "title": "保修政策摘录"},
        ]
        return self.db.kb2_create_batch(TENANT, self.admin["id"], entries)

    def test_create_batch_real_parse_drafts_only_and_persistence(self):
        batch_id = self.create_mixed_batch()
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        self.assertEqual(batch["status"], "draft")
        pdf = self.source_by(batch, "保修说明.pdf")
        self.assertEqual(pdf["status"], "extracted")
        self.assertEqual(pdf["paragraphs"][1]["location"]["page"], 2)
        self.assertTrue(pdf["item_id"])
        docx_source = self.source_by(batch, "安装指南.docx")
        self.assertEqual(docx_source["status"], "extracted")
        self.assertTrue(any(p["source_kind"] == "docx_table" for p in docx_source["paragraphs"]))
        link = self.source_by(batch, "https://example.invalid/manual")
        self.assertEqual(link["status"], "registered")
        self.assertEqual(link["paragraphs"], [])
        manual = self.source_by(batch, "保修政策摘录")
        self.assertEqual(manual["status"], "extracted")
        self.assertEqual(manual["paragraphs"][0]["source_kind"], "manual")
        # 完成识别只创建草稿：绝不自动发布。
        for source in batch["sources"]:
            self.assertTrue(source["item_id"])
            self.assertEqual(source["item_status"], "draft")
        published = self.db.conn.execute(
            "SELECT COUNT(*) FROM knowledge_items WHERE tenant_id=? AND status='published'", (TENANT,)).fetchone()[0]
        self.assertEqual(published, 0)
        # 原件私有 token 命名 + sha256；真实文件落盘且可由本企业读取。
        row = self.db.conn.execute(
            "SELECT * FROM kb2_sources WHERE id=?", (pdf["id"],)).fetchone()
        basename = Path(row["storage_path"]).name
        self.assertRegex(basename, r"^kb2-[0-9a-f]{32}\.pdf$")
        self.assertNotIn("保修说明", str(row["storage_path"]))
        stored = Path(row["storage_path"]).read_bytes()
        import hashlib
        self.assertEqual(row["sha256"], hashlib.sha256(stored).hexdigest())
        sources = self.db.list_knowledge_sources(TENANT, pdf["item_id"])
        self.assertEqual(len(sources), 1)
        _, data = self.db.read_knowledge_source_file(TENANT, sources[0]["id"])
        self.assertEqual(data, stored)
        # 重开数据库：批次、来源、段落、条目全部持久。
        self.reopen()
        again = self.db.kb2_get_batch(TENANT, batch_id)
        self.assertEqual(len(again["sources"]), 4)
        self.assertEqual(self.source_by(again, "保修说明.pdf")["paragraphs"][0]["text"].strip(),
                         "Warranty 36 months.")
        # 跨企业不可见。
        self.assertIsNone(self.db.kb2_get_batch(OTHER, batch_id))

    def test_batch_validation_limits(self):
        with self.assertRaises(ValueError):
            self.db.kb2_create_batch(TENANT, self.admin["id"], [])
        many = [{"kind": "manual", "text": f"条目 {i}", "title": f"条目{i}"} for i in range(21)]
        with self.assertRaises(ValueError):
            self.db.kb2_create_batch(TENANT, self.admin["id"], many)
        self.assertEqual(self.db.conn.execute("SELECT COUNT(*) FROM kb2_batches").fetchone()[0], 0)
        with self.assertRaises(ValueError):
            self.db.kb2_create_batch(TENANT, self.admin["id"], [self.file_entry("程序.exe", b"MZ")])
        with self.assertRaises(ValueError):  # 文本超过 1MB
            self.db.kb2_create_batch(TENANT, self.admin["id"], [self.file_entry("大.txt", b"a" * (1024 * 1024 + 1))])
        with self.assertRaises(ValueError):  # 文档超过 10MB
            self.db.kb2_create_batch(TENANT, self.admin["id"], [self.file_entry("大.pdf", b"%PDF-1.4" + b"0" * (10 * 1024 * 1024))])
        with self.assertRaises(ValueError):  # 空文件
            self.db.kb2_create_batch(TENANT, self.admin["id"], [self.file_entry("空.txt", b"")])
        with self.assertRaises(ValueError):  # 缺链接
            self.db.kb2_create_batch(TENANT, self.admin["id"], [{"kind": "link", "url": ""}])
        with self.assertRaises(ValueError):  # 缺手动文字
            self.db.kb2_create_batch(TENANT, self.admin["id"], [{"kind": "manual", "text": ""}])
        self.assertEqual(self.db.conn.execute("SELECT COUNT(*) FROM kb2_batches").fetchone()[0], 0)
        # 20 件技术上允许。
        twenty = [{"kind": "manual", "text": f"条目 {i}"} for i in range(20)]
        batch_id = self.db.kb2_create_batch(TENANT, self.admin["id"], twenty)
        self.assertEqual(len(self.db.kb2_get_batch(TENANT, batch_id)["sources"]), 20)

    def test_partial_failure_others_continue_and_crash_recovery(self):
        """文件损坏只影响该来源；其他来源继续。崩溃中断（BaseException）后可从数据库恢复。"""
        entries = [self.file_entry("损坏.pdf", b"%PDF-1.4\nbroken"), self.file_entry("正常.txt", "正常内容。".encode())]
        batch_id = self.db.kb2_create_batch(TENANT, self.admin["id"], entries)
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        bad = self.source_by(batch, "损坏.pdf")
        self.assertEqual(bad["status"], "failed")
        self.assertTrue(bad["error"])
        self.assertEqual(bad["item_id"], None)
        good = self.source_by(batch, "正常.txt")
        self.assertEqual(good["status"], "extracted")
        # 崩溃模拟：处理中断（不被 except Exception 捕获），来源留在 processing。
        class Crash(BaseException):
            pass

        original = knowledge_ingestion.extract_document
        knowledge_ingestion.extract_document = lambda *a, **k: (_ for _ in ()).throw(Crash())
        try:
            with self.assertRaises(Crash):
                self.db.kb2_create_batch(TENANT, self.admin["id"], [self.file_entry("指南.pdf", build_text_pdf(["Recover me."]))])
        finally:
            knowledge_ingestion.extract_document = original
        stuck = self.db.conn.execute(
            "SELECT COUNT(*) FROM kb2_sources WHERE tenant_id=? AND status='processing'", (TENANT,)).fetchone()[0]
        self.assertEqual(stuck, 1)
        # 恢复：待处理来源从数据库继续；重复处理不重复创建条目。
        result = self.db.kb2_process_pending_sources(TENANT)
        self.assertGreaterEqual(result["processed"], 1)
        recovered = self.db.conn.execute(
            "SELECT * FROM kb2_sources WHERE tenant_id=? AND filename='指南.pdf'", (TENANT,)).fetchone()
        self.assertEqual(recovered["status"], "extracted")
        items = self.db.conn.execute(
            "SELECT COUNT(*) FROM knowledge_items WHERE tenant_id=? AND id=?", (TENANT, recovered["item_id"])).fetchone()[0]
        self.assertEqual(items, 1)
        self.db.kb2_process_pending_sources(TENANT)  # 再跑一次：无待处理，不重复建
        again = self.db.conn.execute("SELECT COUNT(*) FROM kb2_sources WHERE tenant_id=?", (TENANT,)).fetchone()[0]
        self.assertEqual(again, 3)

    def test_failed_source_retry_idempotent(self):
        """处理暂时失败 → 重试同一来源任务；不重复创建知识，成功后不可重复重试。"""
        original = knowledge_ingestion.extract_document
        knowledge_ingestion.extract_document = lambda *a, **k: {
            "status": "failed", "format": "txt", "paragraphs": [], "warnings": [],
            "error": "临时失败：处理服务不可达", "title_hint": ""}
        try:
            batch_id = self.db.kb2_create_batch(
                TENANT, self.admin["id"], [self.file_entry("重试.txt", "重试内容。".encode())])
        finally:
            knowledge_ingestion.extract_document = original
        source = self.source_by(self.db.kb2_get_batch(TENANT, batch_id), "重试.txt")
        self.assertEqual(source["status"], "failed")
        self.assertIsNone(source["item_id"])
        self.db.kb2_retry_source(TENANT, self.admin["id"], source["id"])
        retried = self.source_by(self.db.kb2_get_batch(TENANT, batch_id), "重试.txt")
        self.assertEqual(retried["status"], "extracted")
        self.assertTrue(retried["item_id"])
        count = self.db.conn.execute(
            "SELECT COUNT(*) FROM knowledge_items WHERE tenant_id=?", (TENANT,)).fetchone()[0]
        self.assertEqual(count, 1)
        with self.assertRaises(ValueError):  # 已成功处理的来源不能再重试
            self.db.kb2_retry_source(TENANT, self.admin["id"], source["id"])

    def test_paragraph_check_edit_exclude_gate_and_publish(self):
        batch_id = self.db.kb2_create_batch(
            TENANT, self.admin["id"], [self.file_entry("条款.txt", "第一条 保修 24 个月。\n第二条 响应 48 小时。\n第三条 内部备注：仅供内部培训。".encode())])
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        source = self.source_by(batch, "条款.txt")
        paragraphs = source["paragraphs"]
        self.assertEqual(len(paragraphs), 3)
        # 未核对：发布被拦。
        self.db.kb2_update_paragraph(TENANT, self.admin["id"], paragraphs[0]["id"], checked=True)
        summary = self.db.kb2_publish_batch(TENANT, self.admin["id"], batch_id)
        self.assertEqual(summary["published"], [])
        self.assertTrue(summary["failed"])
        self.assertIn("未核对", summary["failed"][0]["error"])
        # 编辑第二段：撤销核对并标记人工修改；第三段排除。
        self.db.kb2_update_paragraph(TENANT, self.admin["id"], paragraphs[1]["id"], text="第二条 响应 24 小时。")
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        edited = self.source_by(batch, "条款.txt")["paragraphs"][1]
        self.assertEqual(edited["origin"], "edited")
        self.assertFalse(edited["checked"])
        self.assertEqual(edited["text"], "第二条 响应 24 小时。")
        self.db.kb2_update_paragraph(TENANT, self.admin["id"], paragraphs[1]["id"], checked=True)
        self.db.kb2_update_paragraph(TENANT, self.admin["id"], paragraphs[2]["id"], excluded=True)
        # 保存设置（草稿）不发布。
        self.db.kb2_save_batch_settings(TENANT, self.admin["id"], batch_id, usage="internal", ai_authorized=False)
        self.assertEqual(self.source_by(
            self.db.kb2_get_batch(TENANT, batch_id), "条款.txt")["item_status"], "draft")
        # 发布：正文用修改后的第二段，排除段不进入正文。
        summary = self.db.kb2_publish_batch(TENANT, self.admin["id"], batch_id)
        self.assertEqual(len(summary["published"]), 1)
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        source = self.source_by(batch, "条款.txt")
        self.assertEqual(source["item_status"], "published")
        version = self.db.get_published_knowledge_version(TENANT, source["item_id"])
        content = json.loads(version["content_json"])
        self.assertIn("第二条 响应 24 小时。", content["body"])
        self.assertNotIn("内部备注", content["body"])
        # 发布幂等：再发一次不产生新版本。
        before = len(self.db.list_knowledge_versions(TENANT, source["item_id"]))
        again = self.db.kb2_publish_batch(TENANT, self.admin["id"], batch_id)
        after = len(self.db.list_knowledge_versions(TENANT, source["item_id"]))
        self.assertEqual(before, after)
        self.assertTrue(again["skipped"])

    def test_registered_only_archive_publish_ai_unavailable(self):
        """仅登记原件（链接 / 图片）可归档发布，但 AI 不可用、不算发布知识正文。"""
        entries = [
            {"kind": "link", "url": "https://example.invalid/spec", "title": "在线规格"},
            self.file_entry("照片.png", b"\x89PNG\r\n\x1a\n" + b"0" * 64),
        ]
        batch_id = self.db.kb2_create_batch(TENANT, self.admin["id"], entries)
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        for source in batch["sources"]:
            self.assertEqual(source["status"], "registered")
            self.assertEqual(source["identification"], "registered")
        summary = self.db.kb2_publish_batch(TENANT, self.admin["id"], batch_id)
        self.assertEqual(len(summary["archived"]), 2)
        self.assertEqual(summary["published"], [])
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        for source in batch["sources"]:
            self.assertEqual(source["item_status"], "published")
            item = self.db.get_knowledge_item(TENANT, source["item_id"])
            version = self.db.conn.execute(
                "SELECT * FROM knowledge_versions WHERE id=?", (item["published_version_id"],)).fetchone()
            content = json.loads(version["content_json"])
            self.assertTrue(content.get("kb2_registered_only"))
            self.assertEqual(content.get("body", ""), "")
            self.assertEqual(self.db.list_ai_answerable_knowledge(TENANT, source["item_id"]), [])
        # 归档发布也不允许跳过正文核对：登记来源补充人工正文后仍须核对。
        source = batch["sources"][0]
        self.db.kb2_add_manual_paragraph(TENANT, self.admin["id"], source["id"], "人工整理的规格摘要。")
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        refreshed = self.source_by(batch, source["filename"])
        self.assertEqual(refreshed["identification"], "manual")
        item = self.db.get_knowledge_item(TENANT, source["item_id"])
        self.assertIsNotNone(item["draft_version_id"])  # 已发布资料补人工正文 → 新草稿
        summary = self.db.kb2_publish_batch(TENANT, self.admin["id"], batch_id)
        failed_ids = {entry["source_id"] for entry in summary["failed"]}
        self.assertIn(source["id"], failed_ids)
        manual_paragraph = [p for p in refreshed["paragraphs"] if p["source_kind"] == "manual"][0]
        self.db.kb2_update_paragraph(TENANT, self.admin["id"], manual_paragraph["id"], checked=True)
        summary = self.db.kb2_publish_batch(TENANT, self.admin["id"], batch_id)
        self.assertEqual(len(summary["published"]), 1)

    def test_source_field_provenance(self):
        batch_id = self.db.kb2_create_batch(
            TENANT, self.admin["id"], [self.file_entry("参数.txt", "响应时间 24 小时。".encode())])
        source = self.source_by(self.db.kb2_get_batch(TENANT, batch_id), "参数.txt")
        paragraph = source["paragraphs"][0]
        field_id = self.db.kb2_set_source_field(
            TENANT, self.admin["id"], source["id"],
            label="响应时间", value="24", unit="小时", paragraph_id=paragraph["id"], checked=True)
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        field = self.source_by(batch, "参数.txt")["fields"][0]
        self.assertEqual(field["id"], field_id)
        self.assertEqual((field["label"], field["value"], field["unit"]), ("响应时间", "24", "小时"))
        self.assertEqual(field["paragraph_id"], paragraph["id"])
        self.assertEqual(field["paragraph_location"]["kind"], "line")
        self.assertTrue(field["checked"])
        # 引用别的来源段落 → 拒绝；跨企业 → 拒绝。
        other_batch = self.db.kb2_create_batch(
            TENANT, self.admin["id"], [self.file_entry("其它.txt", "另一段文字。".encode())])
        other_source = self.source_by(self.db.kb2_get_batch(TENANT, other_batch), "其它.txt")
        with self.assertRaises(ValueError):
            self.db.kb2_set_source_field(
                TENANT, self.admin["id"], other_source["id"],
                label="错误来源", value="1", paragraph_id=paragraph["id"])

    def test_permissions_and_cross_tenant(self):
        batch_id = self.db.kb2_create_batch(TENANT, self.admin["id"], [{"kind": "manual", "text": "内容。"}])
        source = self.db.kb2_get_batch(TENANT, batch_id)["sources"][0]
        paragraph_id = source["paragraphs"][0]["id"]
        with self.assertRaises(PermissionError):
            self.db.kb2_create_batch(TENANT, self.member["id"], [{"kind": "manual", "text": "越权。"}])
        with self.assertRaises(PermissionError):
            self.db.kb2_update_paragraph(TENANT, self.member["id"], paragraph_id, checked=True)
        with self.assertRaises(PermissionError):
            self.db.kb2_publish_batch(TENANT, self.member["id"], batch_id)
        with self.assertRaises(PermissionError):
            self.db.kb2_confirm_batch_topics(TENANT, self.member["id"], batch_id)
        with self.assertRaises(PermissionError):
            self.db.kb2_retry_source(TENANT, self.member["id"], source["id"])
        # 对方企业管理员：不能读批次 / 段落 / 重试。
        self.assertIsNone(self.db.kb2_get_batch(OTHER, batch_id))
        with self.assertRaises(PermissionError):
            self.db.kb2_update_paragraph(TENANT, self.other_admin_id, paragraph_id, checked=True)
        with self.assertRaises(ValueError):
            self.db.kb2_retry_source(OTHER, self.other_admin_id, source["id"])
        # 只读成员仍可读。
        self.assertIsNotNone(self.db.kb2_get_batch(TENANT, batch_id))

    def test_topic_suggestion_flow_and_confirm_idempotent(self):
        topic_id = self.db.kb2_create_topic(TENANT, self.admin["id"], "保修政策")
        batch_id = self.db.kb2_create_batch(TENANT, self.admin["id"], [
            {"kind": "manual", "text": "整机保修政策适用于全部机型。\n售后响应 48 小时。", "title": "售后摘录"},
        ])
        source_id = self.db.kb2_get_batch(TENANT, batch_id)["sources"][0]["id"]
        # 系统建议（真实文字命中）：现有主题「保修政策」出现在正文里。
        suggestions = self.db.kb2_build_topic_suggestions(TENANT, self.admin["id"], batch_id)
        text_matches = [s for s in suggestions if s["basis"] == "text_match"]
        self.assertEqual(len(text_matches), 1)
        self.assertEqual(text_matches[0]["existing_topic_id"], topic_id)
        # 任务3 补齐（2026-10-12 用户要求）：资料实际标题也是新主题候选（basis=source_title），
        # 只是候选、仍需成员确认；成员显式输入的建议仍是 basis=member。
        title_candidates = [s for s in suggestions if s["basis"] == "source_title"]
        self.assertEqual([s["name"] for s in title_candidates], ["售后摘录"])
        self.assertIsNone(title_candidates[0]["existing_topic_id"])
        # 整批共同主题 + 改名 + 合并 + 拒绝。
        new_id = self.db.kb2_add_topic_suggestion(TENANT, self.admin["id"], batch_id, "安装与维护")
        self.db.kb2_rename_topic_suggestion(TENANT, self.admin["id"], new_id, "安装与维护流程")
        reject_id = self.db.kb2_add_topic_suggestion(TENANT, self.admin["id"], batch_id, "待拒绝主题")
        self.db.kb2_reject_topic_suggestion(TENANT, self.admin["id"], reject_id)
        merge_id = self.db.kb2_add_topic_suggestion(TENANT, self.admin["id"], batch_id, "售后支持")
        self.db.kb2_merge_topic_suggestion(TENANT, self.admin["id"], merge_id, topic_id)
        # 同批同名合并：再次添加同名 → 仍是同一条建议。
        again_id = self.db.kb2_add_topic_suggestion(TENANT, self.admin["id"], batch_id, "安装与维护流程")
        self.assertEqual(again_id, new_id)
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        self.assertEqual(len([s for s in batch["suggestions"] if s["status"] == "pending"]), 4)
        # 确认前不建立：没有新主题、没有关联。
        self.assertEqual([t["title"] for t in self.db.kb2_list_topics(TENANT)], ["保修政策"])
        self.assertEqual(self.db.kb2_list_relations(TENANT, topic_id), [])
        summary = self.db.kb2_confirm_batch_topics(TENANT, self.admin["id"], batch_id)
        self.assertEqual({t["title"] for t in summary["topics_created"]}, {"安装与维护流程", "售后摘录"})
        topics = {t["title"]: t["id"] for t in self.db.kb2_list_topics(TENANT)}
        self.assertEqual(set(topics), {"保修政策", "安装与维护流程", "售后摘录"})
        self.assertTrue(topics["安装与维护流程"] != topic_id)
        relations = self.db.kb2_list_relations(TENANT, topic_id)
        self.assertEqual(len(relations), 1)  # 两条指向同一主题的建议合并为一条关联
        self.assertEqual(len(self.db.kb2_list_relations(TENANT, topics["安装与维护流程"])), 1)
        self.assertEqual(len(self.db.kb2_list_relations(TENANT, topics["售后摘录"])), 1)
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        statuses = {s["name"]: s["status"] for s in batch["suggestions"]}
        self.assertEqual(statuses["待拒绝主题"], "rejected")
        # 幂等：再次确认不重复建主题 / 关联；被拒绝的建议不会复活。
        self.db.kb2_build_topic_suggestions(TENANT, self.admin["id"], batch_id)
        summary2 = self.db.kb2_confirm_batch_topics(TENANT, self.admin["id"], batch_id)
        self.assertEqual(summary2["topics_created"], [])
        self.assertEqual(len(self.db.kb2_list_topics(TENANT)), 3)
        self.assertEqual(len(self.db.kb2_list_relations(TENANT, topic_id)), 1)
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        self.assertEqual(batch["status"], "confirmed")
        self.assertEqual({s["name"] for s in batch["suggestions"] if s["status"] == "pending"}, set())
        # 批次内资料出现在「全部资料」，并带已确认主题。
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        item_ids = {s["item_id"] for s in batch["sources"] if s["item_id"]}
        self.assertTrue(item_ids)
        materials = {m["id"]: m for m in self.db.kb2_list_materials(TENANT) if m["kind"] == "knowledge"}
        for item_id in item_ids:
            self.assertEqual({t["title"] for t in materials[item_id]["topics"]},
                             {"保修政策", "安装与维护流程", "售后摘录"})


if __name__ == "__main__":
    unittest.main()
