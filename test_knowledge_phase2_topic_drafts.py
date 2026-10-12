"""企业知识第二批 · 任务3 后端补齐：真实主题初稿 / 新主题候选 / 批次找回 / 主题停用。

对应用户 2026-10-12 对任务3的明确要求（补齐 Task3 所需的真实主题初稿）：
- 来源摘录编排：只使用**已核对且未排除**的文字段落，逐段抄录原文与确切出处；
- 确切版本来源钉住：出处固定为资料身份 + 钉住版本 + 原件 + 位置（资料更新不静默换版）；
- 待核对：初稿段落一律「待核对」，核对后才能发布主题；
- 不覆盖人工文本：只追加、不覆盖；同出处同文字不重复追加；
- 系统新主题建议依据**实际标题 / 标题段落**生成（basis=source_title / source_heading），
  一律只是候选、必须人工确认；不声称模型摘要或读取扫描/图片；成员输入仍记为 member；
- 按已批准 PRD 保留基础主题关联与发布历史 / 停用，不扩图谱；
- 批次清单（正式界面找回已保存批次）。

运行环境：系统 python3 可运行；DOCX 标题段落用例需隔离环境
.venv-knowledge（python-docx），在该环境为 0 skip。只用临时数据库与测试私有上传目录；
不触碰 data/ 下真实演示数据与任何密钥。
"""
import io
import json
import os
import tempfile
import unittest
from pathlib import Path

from app import Database

try:
    import pypdf  # noqa: F401
    import docx as docx_lib

    HAVE_PARSERS = True
except ImportError:  # pragma: no cover - 环境说明
    HAVE_PARSERS = False

REQUIRES_PARSERS = unittest.skipUnless(
    HAVE_PARSERS, "需要隔离环境 .venv-knowledge（python-docx）")

TENANT = "demo-tenant"


def build_docx(title_heading="安装指南"):
    document = docx_lib.Document()
    document.add_heading(title_heading, level=1)
    document.add_paragraph("连接电源并检查指示灯。")
    document.add_heading("维护", level=2)
    document.add_paragraph("每月检查滤芯。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TopicDraftTestBase(unittest.TestCase):
    maxDiff = None

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.upload_dir = Path(self.tmp.name) / "uploads"
        os.environ["B2B_KB_STORAGE_ROOT"] = str(self.upload_dir)
        self.db_path = Path(self.tmp.name) / "knowledge2-drafts.db"
        self.db = Database(self.db_path)
        self.db.knowledge_storage_root = self.upload_dir
        self.db.seed_demo_users()
        self.admin = self.db.authenticate_user("demo-admin@example.invalid", "demo-admin")

    def tearDown(self):
        self.db.conn.close()
        os.environ.pop("B2B_KB_STORAGE_ROOT", None)
        self.tmp.cleanup()

    def file_entry(self, filename, data, title=""):
        return {"kind": "file", "filename": filename, "data": data, "title": title}

    def source_by(self, batch, filename):
        return next(source for source in batch["sources"] if source["filename"] == filename)


class TopicDraftBuildTests(TopicDraftTestBase):
    def test_build_topic_draft_only_checked_pinned_and_idempotent(self):
        """只收录已核对未排除段落；出处钉住确切版本与位置；重复整理不重复追加。"""
        batch_id = self.db.kb2_create_batch(
            TENANT, self.admin["id"], [self.file_entry("摘录.txt", "第一行。\n第二行。\n第三行。".encode())])
        source = self.source_by(self.db.kb2_get_batch(TENANT, batch_id), "摘录.txt")
        first, second, third = source["paragraphs"]
        self.db.kb2_update_paragraph(TENANT, self.admin["id"], first["id"], checked=True)
        self.db.kb2_update_paragraph(TENANT, self.admin["id"], third["id"], excluded=True)
        item = self.db.get_knowledge_item(TENANT, source["item_id"])

        topic_id = self.db.kb2_create_topic(TENANT, self.admin["id"], "摘录主题")
        self.assertEqual(self.db.kb2_build_topic_draft(TENANT, self.admin["id"], topic_id)["total"], 0)
        self.db.kb2_set_relation(TENANT, self.admin["id"], topic_id, {"kind": "knowledge", "id": item["id"]})
        result = self.db.kb2_build_topic_draft(TENANT, self.admin["id"], topic_id)
        self.assertEqual(result, {"added": 1, "skipped": 0, "total": 1, "unattributed": 0})

        topic = self.db.kb2_get_topic(TENANT, topic_id)
        paragraph = topic["content"]["paragraphs"][0]
        self.assertEqual(paragraph["text"], "第一行。")
        self.assertFalse(paragraph["checked"], "初稿段落必须默认待核对")
        self.assertEqual(paragraph["origin"], "from_material")
        self.assertEqual(paragraph["source"]["item_id"], item["id"])
        self.assertEqual(paragraph["source"]["version_id"], item["draft_version_id"])
        self.assertEqual(paragraph["source"]["location"], {"kind": "line", "line": 1})
        self.assertIn("第 1 行", paragraph["source_label"])
        self.assertFalse(paragraph["source_stale"])
        # 未核对不能发布主题（初稿默认待核对）。
        with self.assertRaisesRegex(ValueError, "未核对"):
            self.db.kb2_publish_topic_draft(TENANT, self.admin["id"], topic_id)
        # 未核对 / 被排除段落不进入初稿。
        texts = [p["text"] for p in topic["content"]["paragraphs"]]
        self.assertEqual(texts, ["第一行。"])
        # 重复整理：同出处同文字不重复追加（幂等）。
        again = self.db.kb2_build_topic_draft(TENANT, self.admin["id"], topic_id)
        self.assertEqual(again["added"], 0)
        self.assertEqual(again["skipped"], 1)

    def test_build_topic_draft_appends_without_overwriting_manual_and_publish_gate(self):
        """只追加：人工段落原样保留；新段落待核对；全部核对后可发布主题。"""
        topic_id = self.db.kb2_create_topic(TENANT, self.admin["id"], "混合主题")
        fresh = json.loads(self.db._kb2_topic_version_row(
            TENANT, topic_id, draft=True)["content_json"])  # 读取当前草稿（空）
        self.assertEqual(fresh.get("paragraphs"), [])
        self.db.kb2_save_topic_draft(
            TENANT, self.admin["id"], topic_id,
            paragraphs=[{"id": "manual-1", "kind": "intro", "text": "成员撰写的主题介绍。",
                         "origin": "manual", "checked": True, "source": {"kind": "manual"}}])

        batch_id = self.db.kb2_create_batch(
            TENANT, self.admin["id"], [self.file_entry("补充.txt", "资料里的核对段落。".encode())])
        source = self.source_by(self.db.kb2_get_batch(TENANT, batch_id), "补充.txt")
        self.db.kb2_update_paragraph(TENANT, self.admin["id"], source["paragraphs"][0]["id"], checked=True)
        self.db.kb2_set_relation(
            TENANT, self.admin["id"], topic_id, {"kind": "knowledge", "id": source["item_id"]})

        self.db.kb2_build_topic_draft(TENANT, self.admin["id"], topic_id)
        topic = self.db.kb2_get_topic(TENANT, topic_id)
        paragraphs = topic["content"]["paragraphs"]
        self.assertEqual([p["id"] for p in paragraphs][0], "manual-1")
        self.assertEqual(paragraphs[0]["text"], "成员撰写的主题介绍。")
        self.assertTrue(paragraphs[0]["checked"], "人工段落原样保留（核对状态不被覆盖）")
        self.assertEqual(paragraphs[0]["source_label"], "人工撰写")
        self.assertEqual(len(paragraphs), 2)
        self.assertEqual(paragraphs[1]["text"], "资料里的核对段落。")
        # 仍有一段未核对：发布拦截并定位段号；核对后发布成功。
        with self.assertRaisesRegex(ValueError, "第 2 段"):
            self.db.kb2_publish_topic_draft(TENANT, self.admin["id"], topic_id)
        self.db.kb2_save_topic_draft(
            TENANT, self.admin["id"], topic_id,
            paragraphs=[dict(paragraphs[0]), {**paragraphs[1], "checked": True}])
        version = self.db.kb2_publish_topic_draft(TENANT, self.admin["id"], topic_id)
        self.assertEqual(version, 1)
        refreshed = self.db.kb2_get_topic(TENANT, topic_id)
        self.assertEqual(refreshed["status"], "published")
        self.assertEqual(len(refreshed["versions"]), 1)


class NewTopicCandidateTests(TopicDraftTestBase):
    @REQUIRES_PARSERS
    def test_candidates_from_title_and_headings_are_candidates_until_confirmed(self):
        """新主题候选来自实际标题与标题段落：确认前不建立；拒绝不复活；失败资料不产出候选。"""
        batch_id = self.db.kb2_create_batch(
            TENANT, self.admin["id"], [
                self.file_entry("设备安装指南.docx", build_docx()),
                self.file_entry("坏标题.txt", b"not text \x00 binary"),
            ])
        batch = self.db.kb2_get_batch(TENANT, batch_id)
        self.assertEqual(self.source_by(batch, "坏标题.txt")["status"], "failed")
        suggestions = self.db.kb2_build_topic_suggestions(TENANT, self.admin["id"], batch_id)
        by_name = {s["name"]: s for s in suggestions}
        self.assertEqual(by_name["设备安装指南"]["basis"], "source_title")
        self.assertEqual(by_name["安装指南"]["basis"], "source_heading")
        self.assertEqual(by_name["维护"]["basis"], "source_heading")
        self.assertNotIn("坏标题", by_name, "解析失败的资料不产出标题候选")
        for suggestion in suggestions:
            self.assertIsNone(suggestion["existing_topic_id"])
            self.assertEqual(suggestion["status"], "pending")
        # 确认前不建立任何主题。
        self.assertEqual(self.db.kb2_list_topics(TENANT), [])
        # 成员输入仍是 member 来源（不冒充系统候选）。
        member_id = self.db.kb2_add_topic_suggestion(TENANT, self.admin["id"], batch_id, "成员添加的主题")
        suggestions = self.db.kb2_build_topic_suggestions(TENANT, self.admin["id"], batch_id)
        member_view = next(s for s in suggestions if s["id"] == member_id)
        self.assertEqual(member_view["basis"], "member")
        # 拒绝的候选不复活、不加资料。
        head_id = by_name["维护"]["id"]
        self.db.kb2_reject_topic_suggestion(TENANT, self.admin["id"], head_id)
        again = {s["name"]: s for s in self.db.kb2_build_topic_suggestions(TENANT, self.admin["id"], batch_id)}
        self.assertEqual(again["维护"]["status"], "rejected")
        # 确认后才建立（草稿状态，不因批次确认变成已发布）。
        self.db.kb2_confirm_batch_topics(TENANT, self.admin["id"], batch_id)
        topics = {t["title"]: t for t in self.db.kb2_list_topics(TENANT)}
        self.assertEqual(set(topics), {"设备安装指南", "安装指南", "成员添加的主题"})
        for topic in topics.values():
            self.assertEqual(topic["status"], "draft")

    def test_candidates_do_not_duplicate_existing_topic_names(self):
        """候选不重复现有主题名（避免同义重复建议）；同名走文字命中合并。"""
        topic_id = self.db.kb2_create_topic(TENANT, self.admin["id"], "保修政策")
        # 手动文字来源：标题即正文首行。
        batch_id = self.db.kb2_create_batch(
            TENANT, self.admin["id"], [
                {"kind": "manual", "text": "保修政策适用于全部机型。", "title": "保修政策"},
            ])
        suggestions = self.db.kb2_build_topic_suggestions(TENANT, self.admin["id"], batch_id)
        named = [s for s in suggestions if s["name"] == "保修政策"]
        self.assertEqual(len(named), 1, "已有同名主题：只保留文字命中，不重复造候选")
        self.assertEqual(named[0]["basis"], "text_match")
        self.assertEqual(named[0]["existing_topic_id"], topic_id)


class TopicLifecycleAndBatchListTests(TopicDraftTestBase):
    def test_topic_disable_enable_keeps_history(self):
        """停用 / 恢复：发布历史与关联保留；草稿不能停用；重复调用幂等。"""
        topic_id = self.db.kb2_create_topic(TENANT, self.admin["id"], "停用主题")
        with self.assertRaises(ValueError):
            self.db.kb2_set_topic_disabled(TENANT, self.admin["id"], topic_id)
        self.db.kb2_save_topic_draft(
            TENANT, self.admin["id"], topic_id,
            paragraphs=[{"kind": "intro", "text": "正文。", "origin": "manual", "checked": True,
                         "source": {"kind": "manual"}}])
        self.assertEqual(self.db.kb2_publish_topic_draft(TENANT, self.admin["id"], topic_id), 1)
        self.db.kb2_set_topic_disabled(TENANT, self.admin["id"], topic_id)
        topic = self.db.kb2_get_topic(TENANT, topic_id)
        self.assertEqual(topic["status"], "disabled")
        self.assertIsNotNone(topic["published_version_id"], "停用保留已发布版本（历史不删除）")
        self.db.kb2_set_topic_disabled(TENANT, self.admin["id"], topic_id)  # 幂等
        self.db.kb2_set_topic_disabled(TENANT, self.admin["id"], topic_id, disabled=False)
        self.assertEqual(self.db.kb2_get_topic(TENANT, topic_id)["status"], "published")
        self.assertEqual(len(self.db.kb2_list_topic_versions(TENANT, topic_id)), 1)

    def test_list_batches_for_recovery_is_read_only(self):
        """批次清单：正式界面可找回已保存批次；只读，不改变状态。"""
        first = self.db.kb2_create_batch(
            TENANT, self.admin["id"], [{"kind": "manual", "text": "甲。", "title": "甲批次"}], process_now=False)
        second = self.db.kb2_create_batch(
            TENANT, self.admin["id"], [
                {"kind": "manual", "text": "乙。", "title": "乙批次"},
                {"kind": "link", "url": "https://example.invalid/乙"},
            ], process_now=False)
        self.db.conn.execute(
            "UPDATE kb2_batches SET created_at='2020-01-01T00:00:00+00:00' WHERE id=?", (first,))
        self.db.conn.commit()
        rows = self.db.kb2_list_batches(TENANT)
        self.assertEqual([row["id"] for row in rows], [second, first])
        by_id = {row["id"]: row for row in rows}
        self.assertEqual(by_id[second]["sources_total"], 2)
        self.assertEqual(by_id[second]["sources_pending"], 2)
        self.assertEqual(by_id[first]["sources_total"], 1)
        self.assertEqual(by_id[first]["status"], "draft")
        # 只读：清单不触发处理。
        statuses = [row["status"] for row in self.db.conn.execute(
            "SELECT status FROM kb2_sources WHERE batch_id=?", (second,)).fetchall()]
        self.assertEqual(statuses, ["pending", "pending"])


class TopicScopeGapTests(TopicDraftTestBase):
    """独立验收「范围核对」补齐：相关主题基础关系 / 发布设置 / 历史恢复（不扩图谱）。"""

    def test_related_topics_links_are_bidirectional_and_independent(self):
        """相关主题（主题↔主题）：双向可见、幂等、只动关系；不改变内容、发布状态与授权。"""
        first = self.db.kb2_create_topic(TENANT, self.admin["id"], "相关甲")
        second = self.db.kb2_create_topic(TENANT, self.admin["id"], "相关乙")
        self.db.kb2_set_topic_link(TENANT, self.admin["id"], first, second, linked=True)
        self.db.kb2_set_topic_link(TENANT, self.admin["id"], first, second, linked=True)  # 幂等
        self.assertEqual(self.db.conn.execute("SELECT COUNT(*) FROM kb2_topic_links").fetchone()[0], 1)
        self.assertEqual([item["id"] for item in self.db.kb2_list_topic_links(TENANT, first)], [second])
        self.assertEqual([item["id"] for item in self.db.kb2_list_topic_links(TENANT, second)], [first],
                         "相关主题双向可见")
        view = self.db.kb2_get_topic(TENANT, first)
        self.assertEqual([item["title"] for item in view["related_topics"]], ["相关乙"])
        # 关系不影响内容/状态/版本。
        self.assertEqual(view["status"], "draft")
        self.assertEqual(view["content"]["paragraphs"], [])
        with self.assertRaisesRegex(ValueError, "另一个主题"):
            self.db.kb2_set_topic_link(TENANT, self.admin["id"], first, first, linked=True)
        # 反向移除同样有效（双向关系一样看待）；重复移除幂等。
        self.db.kb2_set_topic_link(TENANT, self.admin["id"], second, first, linked=False)
        self.assertEqual(self.db.kb2_list_topic_links(TENANT, first), [])
        self.db.kb2_set_topic_link(TENANT, self.admin["id"], second, first, linked=False)

    def test_restore_historical_version_creates_new_draft_only(self):
        """历史恢复：产生新草稿（注明恢复来源），不覆盖历史、不影响线上；已发布线上版本保持。"""
        topic_id = self.db.kb2_create_topic(TENANT, self.admin["id"], "恢复主题")
        self.db.kb2_save_topic_draft(
            TENANT, self.admin["id"], topic_id,
            paragraphs=[{"id": "p1", "kind": "intro", "text": "第一版。", "origin": "manual",
                         "checked": True, "source": {"kind": "manual"}}])
        self.assertEqual(self.db.kb2_publish_topic_draft(TENANT, self.admin["id"], topic_id), 1)
        self.db.kb2_save_topic_draft(
            TENANT, self.admin["id"], topic_id,
            paragraphs=[{"id": "p1", "kind": "intro", "text": "第二版。", "origin": "manual",
                         "checked": False, "source": {"kind": "manual"}}])
        self.db.kb2_save_topic_draft(
            TENANT, self.admin["id"], topic_id,
            paragraphs=[{"id": "p1", "kind": "intro", "text": "第二版。", "origin": "manual",
                         "checked": True, "source": {"kind": "manual"}}])
        self.assertEqual(self.db.kb2_publish_topic_draft(TENANT, self.admin["id"], topic_id), 2)
        published = self.db.kb2_get_topic(TENANT, topic_id)
        self.assertEqual(published["published_version_number"], 2)
        # 恢复 v1 → 新草稿 v3（restored_from_version=1），线上仍是 v2。
        draft_version_id = self.db.kb2_restore_topic_version(TENANT, self.admin["id"], topic_id, 1)
        topic = self.db.kb2_get_topic(TENANT, topic_id)
        self.assertEqual(topic["draft_version_id"], draft_version_id)
        self.assertEqual(topic["published_version_number"], 2, "恢复不影响线上版本")
        self.assertEqual(topic["content"]["paragraphs"][0]["text"], "第一版。")
        versions = self.db.kb2_list_topic_versions(TENANT, topic_id)
        restored = next(v for v in versions if v["id"] == draft_version_id)
        self.assertEqual(restored["restored_from_version"], 1)
        # 已有草稿时不能重复恢复。
        with self.assertRaisesRegex(ValueError, "已有未发布草稿"):
            self.db.kb2_restore_topic_version(TENANT, self.admin["id"], topic_id, 1)

    def test_topic_publish_settings_roundtrip(self):
        """主题发布设置（用途/AI/到期）走草稿、发布时生效；对外+AI 需要一次确认。"""
        topic_id = self.db.kb2_create_topic(TENANT, self.admin["id"], "设置主题")
        self.db.kb2_save_topic_draft(
            TENANT, self.admin["id"], topic_id,
            paragraphs=[{"id": "p1", "kind": "intro", "text": "对外正文。", "origin": "manual",
                         "checked": True, "source": {"kind": "manual"}}],
            usage="external", ai_authorized=True, expires_at="2030-01-01T12:00")
        topic = self.db.kb2_get_topic(TENANT, topic_id)
        self.assertEqual(topic["usage"], "external")
        self.assertTrue(topic["ai_authorized"])
        self.assertTrue(str(topic["expires_at"]).startswith("2030-01-01"))
        with self.assertRaisesRegex(ValueError, "确认"):
            self.db.kb2_publish_topic_draft(TENANT, self.admin["id"], topic_id)
        self.db.kb2_publish_topic_draft(TENANT, self.admin["id"], topic_id, confirm_ai=True)
        published = self.db.kb2_get_topic(TENANT, topic_id)
        self.assertEqual(published["usage"], "external")
        self.assertTrue(published["ai_authorized"])


if __name__ == "__main__":
    unittest.main()
